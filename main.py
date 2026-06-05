#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
异环 NTE 全角色养成计算器 - 主入口

使用方法:
    python main.py

生成的 Excel 文件将保存在 output/ 目录下。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from config import OUTPUT_DIR, OUTPUT_FILENAME
from generators.calculator_generator import CalculatorGenerator
from generators.character_sheet_generator import CharacterSheetGenerator
from generators.arc_sheet_generator import ArcSheetGenerator
from generators.other_sheets_generator import OtherSheetsGenerator


def ensure_output_dir():
    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR)
    return os.path.join(OUTPUT_DIR, OUTPUT_FILENAME)


def generate_workbook():
    print("=" * 60)
    print("异环 NTE 全角色养成计算器 - Excel 生成工具")
    print("=" * 60)

    wb = Workbook()
    
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    print("[1/5] 生成养成计算器主表...")
    calc_gen = CalculatorGenerator(wb)
    calc_gen.generate()

    print("[2/5] 生成角色等级突破材料总表...")
    char_gen = CharacterSheetGenerator(wb)
    char_gen.generate()

    print("[3/5] 生成弧盘强化总表...")
    arc_gen = ArcSheetGenerator(wb)
    arc_gen.generate()

    print("[4/5] 生成其他副表...")
    other_gen = OtherSheetsGenerator(wb)
    other_gen.generate_all()

    output_path = ensure_output_dir()
    print("[5/5] 保存文件到: %s" % output_path)
    wb.save(output_path)

    print("=" * 60)
    print("生成完成! 文件路径: %s" % output_path)
    print("=" * 60)
    print()
    print("工作表清单:")
    for i, sheet_name in enumerate(wb.sheetnames, 1):
        print("  %d. %s" % (i, sheet_name))
    print()
    print("使用提示:")
    print("  - 黄色区域: 录入区(手动填写角色名、等级、弧盘名)")
    print("  - 浅蓝区域: 名称匹配区(自动 VLOOKUP 匹配材料名称)")
    print("  - 绿色区域: 计算区(自动 SUMIF 计算缺口数量)")
    print("  - 橙色区域: 汇总区(总金币、总异像素材等)")
    print("  - 红色区域: 全队汇总(10个角色 SUM 汇总)")
    print()
    print("数据维护:")
    print("  - 修改 data/character_data.py 更新角色突破数据")
    print("  - 修改 data/arc_data.py 更新弧盘强化数据")
    print("  - 修改 data/material_data.py 更新材料/合成数据")
    print("  - 修改 config.py 调整颜色、布局等配置")

    return output_path


if __name__ == "__main__":
    generate_workbook()
