"""
弧盘强化总表生成器
"""
from config import Colors
from data.arc_data import ARC_LEVEL_UP_DATA, ARC_LEVEL_COLUMNS
from utils.styles import create_header_style, apply_dict_style
from openpyxl.styles import Font, PatternFill, Alignment

class ArcSheetGenerator:
    """弧盘强化总表生成器"""

    def __init__(self, workbook):
        self.wb = workbook
        self.ws = None

    def generate(self):
        """生成弧盘强化总表"""
        self.ws = self.wb.create_sheet("弧盘强化总表")

        # 标题
        self.ws.merge_cells("A1:M1")
        cell = self.ws["A1"]
        cell.value = "弧盘强化材料总表（累计值）"
        cell.font = Font(size=14, bold=True, color=Colors.TITLE_TEXT)
        cell.fill = PatternFill(start_color=Colors.TITLE_BG, end_color=Colors.TITLE_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        self.ws.row_dimensions[1].height = 26

        # 表头
        header_style = create_header_style(Colors.HEADER_ARC_BG)
        for col, val in enumerate(ARC_LEVEL_COLUMNS, 1):
            cell = self.ws.cell(row=3, column=col, value=val)
            apply_dict_style(cell, header_style)

        # 数据
        row = 4
        for level, data in sorted(ARC_LEVEL_UP_DATA.items()):
            for col, val in enumerate([level] + data, 1):
                cell = self.ws.cell(row=row, column=col, value=val)
                cell.border = header_style["border"]
                cell.alignment = Alignment(horizontal="center", vertical="center")
            self.ws.row_dimensions[row].height = 22
            row += 1

        # 说明
        self.ws.merge_cells(f"A{row}:M{row}")
        cell = self.ws.cell(row=row, column=1)
        cell.value = "注：以上为示例数据，弧盘等级请按游戏实际最大等级调整。"
        cell.font = Font(italic=True, color="999999")
        cell.alignment = Alignment(horizontal="left", vertical="center")

        # 列宽
        for c in range(1, 14):
            from openpyxl.utils import get_column_letter
            self.ws.column_dimensions[get_column_letter(c)].width = 16

        return self.ws
