"""
养成计算器主表生成器
"""
from openpyxl.utils import get_column_letter
from config import Layout, Colors
from utils.styles import (
    create_header_style, create_input_style, create_name_style,
    create_calc_style, create_sum_style, create_total_row_style,
    apply_dict_style
)
from utils.formula_builder import FormulaBuilder

class CalculatorGenerator:
    """养成计算器主表生成器"""

    def __init__(self, workbook):
        self.wb = workbook
        self.ws = None

    def generate(self):
        """生成养成计算器主表"""
        # 创建或获取sheet
        if "养成计算器" in self.wb.sheetnames:
            del self.wb["养成计算器"]
        self.ws = self.wb.create_sheet("养成计算器", 0)

        self._generate_title()
        self._generate_header()
        self._generate_data_rows()
        self._generate_summary_row()
        self._adjust_column_widths()

        return self.ws

    def _generate_title(self):
        """生成标题行"""
        from openpyxl.styles import Font, PatternFill, Alignment

        # 标题
        self.ws.merge_cells("A1:AO1")
        cell = self.ws["A1"]
        cell.value = "异环 NTE 全角色养成计算器（突破 + 弧盘 联动版）—— 录入角色名/弧盘名自动匹配材料"
        cell.font = Font(size=14, bold=True, color=Colors.TITLE_TEXT)
        cell.fill = PatternFill(start_color=Colors.TITLE_BG, end_color=Colors.TITLE_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.ws.row_dimensions[1].height = 32

        # 说明
        self.ws.merge_cells("A2:AO2")
        cell = self.ws["A2"]
        cell.value = "使用说明：黄色区域为录入区；浅蓝色区域为自动匹配的材料名称；绿色区域为自动计算的缺口数量；橙色区域为汇总。请勿手动修改公式。"
        cell.font = Font(italic=True, color="666666", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.ws.row_dimensions[2].height = 28

    def _generate_header(self):
        """生成表头"""
        headers = [
            # 录入区 (A-H)
            ("序号", Colors.HEADER_INPUT_BG), ("角色名", Colors.HEADER_INPUT_BG), 
            ("当前等级", Colors.HEADER_INPUT_BG), ("目标等级", Colors.HEADER_INPUT_BG),
            ("弧盘名", Colors.HEADER_INPUT_BG), ("当前弧盘等级", Colors.HEADER_INPUT_BG),
            ("目标弧盘等级", Colors.HEADER_INPUT_BG), ("备注", Colors.HEADER_INPUT_BG),
            # 角色突破区 (I-T)
            ("突破素材\n名称", Colors.HEADER_CHAR_BG), ("突破素材\n数量", Colors.HEADER_CHAR_BG),
            ("低级异像\n名称", Colors.HEADER_CHAR_BG), ("低级异像\n数量", Colors.HEADER_CHAR_BG),
            ("中级异像\n名称", Colors.HEADER_CHAR_BG), ("中级异像\n数量", Colors.HEADER_CHAR_BG),
            ("高级异像\n名称", Colors.HEADER_CHAR_BG), ("高级异像\n数量", Colors.HEADER_CHAR_BG),
            ("新锐\n攻略", Colors.HEADER_CHAR_BG), ("资深\n攻略", Colors.HEADER_CHAR_BG),
            ("特级\n攻略", Colors.HEADER_CHAR_BG), ("角色突破\n金币", Colors.HEADER_CHAR_BG),
            # 弧盘强化区 (U-AJ)
            ("低级弧盘\n名称", Colors.HEADER_ARC_BG), ("低级弧盘\n数量", Colors.HEADER_ARC_BG),
            ("中级弧盘\n名称", Colors.HEADER_ARC_BG), ("中级弧盘\n数量", Colors.HEADER_ARC_BG),
            ("高级弧盘\n名称", Colors.HEADER_ARC_BG), ("高级弧盘\n数量", Colors.HEADER_ARC_BG),
            ("弧盘低级\n异像名称", Colors.HEADER_ARC_BG), ("弧盘低级\n异像数量", Colors.HEADER_ARC_BG),
            ("弧盘中级\n异像名称", Colors.HEADER_ARC_BG), ("弧盘中级\n异像数量", Colors.HEADER_ARC_BG),
            ("弧盘高级\n异像名称", Colors.HEADER_ARC_BG), ("弧盘高级\n异像数量", Colors.HEADER_ARC_BG),
            ("淡色\n染剂", Colors.HEADER_ARC_BG), ("无彩\n染剂", Colors.HEADER_ARC_BG),
            ("混沌\n染剂", Colors.HEADER_ARC_BG), ("弧盘强化\n金币", Colors.HEADER_ARC_BG),
            # 汇总区 (AK-AO)
            ("总金币", Colors.HEADER_SUM_BG), ("总染剂\n经验", Colors.HEADER_SUM_BG),
            ("总低级\n异像", Colors.HEADER_SUM_BG), ("总中级\n异像", Colors.HEADER_SUM_BG),
            ("总高级\n异像", Colors.HEADER_SUM_BG),
        ]

        for col, (val, bg_color) in enumerate(headers, 1):
            cell = self.ws.cell(row=Layout.HEADER_ROW, column=col, value=val)
            style = create_header_style(bg_color)
            apply_dict_style(cell, style)

        self.ws.row_dimensions[Layout.HEADER_ROW].height = 36

    def _generate_data_rows(self):
        """生成10个角色的数据行"""
        input_style = create_input_style()
        name_style = create_name_style()
        calc_style = create_calc_style()
        sum_style = create_sum_style()

        for i in range(1, Layout.MAX_CHARACTERS + 1):
            data_row = Layout.DATA_START_ROW + i - 1  # 5-14行

            # ---- 录入区 (A-H) ----
            self.ws.cell(row=data_row, column=Layout.COL_SEQ, value=i)
            for c in range(Layout.COL_CHAR_NAME, Layout.COL_REMARK + 1):
                cell = self.ws.cell(row=data_row, column=c, value="")
                apply_dict_style(cell, input_style)

            # ---- 角色突破区 (I-T) ----
            # I: 突破素材名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_BREAK_MAT_NAME)
            cell.value = f'=IF(B{data_row}="","",IFERROR(VLOOKUP(B{data_row},全角色突破材料对照!$A:$E,2,FALSE),"未找到"))'
            apply_dict_style(cell, name_style)

            # J: 突破素材数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_BREAK_MAT_NUM)
            cell.value = FormulaBuilder.character_break_material_num(data_row)
            apply_dict_style(cell, calc_style)

            # K: 低级异像名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_LOW_YIXIANG_NAME)
            cell.value = f'=IF(B{data_row}="","",IFERROR(VLOOKUP(B{data_row},全角色突破材料对照!$A:$E,3,FALSE),""))'
            apply_dict_style(cell, name_style)

            # L: 低级异像数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_LOW_YIXIANG_NUM)
            cell.value = FormulaBuilder.character_low_yixiang_num(data_row)
            apply_dict_style(cell, calc_style)

            # M: 中级异像名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_MID_YIXIANG_NAME)
            cell.value = f'=IF(B{data_row}="","",IFERROR(VLOOKUP(B{data_row},全角色突破材料对照!$A:$E,4,FALSE),""))'
            apply_dict_style(cell, name_style)

            # N: 中级异像数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_MID_YIXIANG_NUM)
            cell.value = FormulaBuilder.character_mid_yixiang_num(data_row)
            apply_dict_style(cell, calc_style)

            # O: 高级异像名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_HIGH_YIXIANG_NAME)
            cell.value = f'=IF(B{data_row}="","",IFERROR(VLOOKUP(B{data_row},全角色突破材料对照!$A:$E,5,FALSE),""))'
            apply_dict_style(cell, name_style)

            # P: 高级异像数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_HIGH_YIXIANG_NUM)
            cell.value = FormulaBuilder.character_high_yixiang_num(data_row)
            apply_dict_style(cell, calc_style)

            # Q-R-S: 经验书数量
            for col_idx, col_name in [(Layout.COL_EXP_BOOK_LOW, "F"), 
                                       (Layout.COL_EXP_BOOK_MID, "G"), 
                                       (Layout.COL_EXP_BOOK_HIGH, "H")]:
                cell = self.ws.cell(row=data_row, column=col_idx)
                cell.value = FormulaBuilder.conditional_formula(
                    f'OR(B{data_row}="",C{data_row}="",D{data_row}="")',
                    '""',
                    FormulaBuilder.sumif_diff("角色等级突破材料总表", "A", f"D{data_row}", f"C{data_row}", col_name)
                )
                apply_dict_style(cell, calc_style)

            # T: 角色突破金币
            cell = self.ws.cell(row=data_row, column=Layout.COL_CHAR_GOLD)
            cell.value = FormulaBuilder.conditional_formula(
                f'OR(B{data_row}="",C{data_row}="",D{data_row}="")',
                '""',
                FormulaBuilder.sumif_diff("角色等级突破材料总表", "A", f"D{data_row}", f"C{data_row}", "I")
            )
            apply_dict_style(cell, calc_style)

            # ---- 弧盘强化区 (U-AJ) ----
            # U: 低级弧盘素材名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_LOW_MAT_NAME)
            cell.value = f'=IF(E{data_row}="","",IFERROR(VLOOKUP(E{data_row},全弧盘突破材料对照!$A:$H,6,FALSE),"未找到"))'
            apply_dict_style(cell, name_style)

            # V: 低级弧盘素材数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_LOW_MAT_NUM)
            cell.value = FormulaBuilder.arc_material_num(data_row, "B")
            apply_dict_style(cell, calc_style)

            # W: 中级弧盘素材名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_MID_MAT_NAME)
            cell.value = f'=IF(E{data_row}="","",IFERROR(VLOOKUP(E{data_row},全弧盘突破材料对照!$A:$H,7,FALSE),""))'
            apply_dict_style(cell, name_style)

            # X: 中级弧盘素材数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_MID_MAT_NUM)
            cell.value = FormulaBuilder.arc_material_num(data_row, "C")
            apply_dict_style(cell, calc_style)

            # Y: 高级弧盘素材名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_HIGH_MAT_NAME)
            cell.value = f'=IF(E{data_row}="","",IFERROR(VLOOKUP(E{data_row},全弧盘突破材料对照!$A:$H,8,FALSE),""))'
            apply_dict_style(cell, name_style)

            # Z: 高级弧盘素材数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_HIGH_MAT_NUM)
            cell.value = FormulaBuilder.arc_material_num(data_row, "D")
            apply_dict_style(cell, calc_style)

            # AA-AF: 弧盘异像素材
            # 弧盘低级异像名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_LOW_YX_NAME)
            cell.value = f'=IF(E{data_row}="","",IFERROR(VLOOKUP(E{data_row},全弧盘突破材料对照!$A:$H,3,FALSE),""))'
            apply_dict_style(cell, name_style)

            # 弧盘低级异像数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_LOW_YX_NUM)
            cell.value = FormulaBuilder.arc_material_num(data_row, "E")
            apply_dict_style(cell, calc_style)

            # 弧盘中级异像名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_MID_YX_NAME)
            cell.value = f'=IF(E{data_row}="","",IFERROR(VLOOKUP(E{data_row},全弧盘突破材料对照!$A:$H,4,FALSE),""))'
            apply_dict_style(cell, name_style)

            # 弧盘中级异像数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_MID_YX_NUM)
            cell.value = FormulaBuilder.arc_material_num(data_row, "F")
            apply_dict_style(cell, calc_style)

            # 弧盘高级异像名称
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_HIGH_YX_NAME)
            cell.value = f'=IF(E{data_row}="","",IFERROR(VLOOKUP(E{data_row},全弧盘突破材料对照!$A:$H,5,FALSE),""))'
            apply_dict_style(cell, name_style)

            # 弧盘高级异像数量
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_HIGH_YX_NUM)
            cell.value = FormulaBuilder.arc_material_num(data_row, "G")
            apply_dict_style(cell, calc_style)

            # AG-AI: 染剂
            for col_idx, col_name in [(Layout.COL_DYE_LIGHT, "H"), 
                                      (Layout.COL_DYE_COLORLESS, "I"), 
                                      (Layout.COL_DYE_CHAOS, "J")]:
                cell = self.ws.cell(row=data_row, column=col_idx)
                cell.value = FormulaBuilder.arc_material_num(data_row, col_name)
                apply_dict_style(cell, calc_style)

            # AJ: 弧盘强化金币
            cell = self.ws.cell(row=data_row, column=Layout.COL_ARC_GOLD)
            cell.value = FormulaBuilder.arc_material_num(data_row, "K")
            apply_dict_style(cell, calc_style)

            # ---- 汇总区 (AK-AO) ----
            # AK: 总金币
            cell = self.ws.cell(row=data_row, column=Layout.COL_TOTAL_GOLD)
            cell.value = FormulaBuilder.total_gold(data_row)
            apply_dict_style(cell, sum_style)

            # AL: 总染剂经验
            cell = self.ws.cell(row=data_row, column=Layout.COL_TOTAL_DYE_EXP)
            cell.value = FormulaBuilder.total_dye_exp(data_row)
            apply_dict_style(cell, sum_style)

            # AM: 总低级异像
            cell = self.ws.cell(row=data_row, column=Layout.COL_TOTAL_LOW_YX)
            cell.value = FormulaBuilder.total_low_yixiang(data_row)
            apply_dict_style(cell, sum_style)

            # AN: 总中级异像
            cell = self.ws.cell(row=data_row, column=Layout.COL_TOTAL_MID_YX)
            cell.value = FormulaBuilder.total_mid_yixiang(data_row)
            apply_dict_style(cell, sum_style)

            # AO: 总高级异像
            cell = self.ws.cell(row=data_row, column=Layout.COL_TOTAL_HIGH_YX)
            cell.value = FormulaBuilder.total_high_yixiang(data_row)
            apply_dict_style(cell, sum_style)

            self.ws.row_dimensions[data_row].height = 22

    def _generate_summary_row(self):
        """生成全队汇总行"""
        total_style = create_total_row_style()
        sum_style = create_sum_style()

        self.ws.merge_cells(f"A{Layout.SUM_ROW}:H{Layout.SUM_ROW}")
        cell = self.ws.cell(row=Layout.SUM_ROW, column=1, value="全队汇总")
        apply_dict_style(cell, total_style)
        self.ws.row_dimensions[Layout.SUM_ROW].height = 26

        for c in range(9, 42):
            col_letter = get_column_letter(c)
            cell = self.ws.cell(row=Layout.SUM_ROW, column=c)
            cell.value = f'=SUM({col_letter}{Layout.DATA_START_ROW}:{col_letter}{Layout.DATA_START_ROW + Layout.MAX_CHARACTERS - 1})'
            apply_dict_style(cell, sum_style)

    def _adjust_column_widths(self):
        """调整列宽"""
        self.ws.column_dimensions['A'].width = 6
        self.ws.column_dimensions['B'].width = 12
        for c in ['C', 'D', 'E', 'F', 'G']:
            self.ws.column_dimensions[c].width = 11
        self.ws.column_dimensions['H'].width = 10

        for c in ['I', 'K', 'M', 'O']:
            self.ws.column_dimensions[c].width = 12
        for c in ['J', 'L', 'N', 'P', 'Q', 'R', 'S', 'T']:
            self.ws.column_dimensions[c].width = 10

        for c in ['U', 'W', 'Y', 'AA', 'AC', 'AE']:
            self.ws.column_dimensions[c].width = 12
        for c in ['V', 'X', 'Z', 'AB', 'AD', 'AF', 'AG', 'AH', 'AI', 'AJ']:
            self.ws.column_dimensions[c].width = 10

        for c in ['AK', 'AL', 'AM', 'AN', 'AO']:
            self.ws.column_dimensions[c].width = 11
