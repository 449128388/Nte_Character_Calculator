"""
异环 NTE 养成计算器 - 全局配置
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==================== 输出配置 ====================
OUTPUT_DIR = "output"
OUTPUT_FILENAME = "异环_NTE_全角色养成计算器_联动版.xlsx"

# ==================== 颜色配置 ====================
class Colors:
    """Excel 颜色常量"""
    TITLE_BG = "203864"
    TITLE_TEXT = "FFFFFF"

    HEADER_INPUT_BG = "F4B084"      # 录入区表头
    HEADER_CHAR_BG = "A9D08E"      # 角色突破区表头
    HEADER_ARC_BG = "9BC2E6"       # 弧盘强化区表头
    HEADER_SUM_BG = "C6E0B4"      # 汇总区表头

    INPUT_BG = "FFF2CC"             # 录入区底色
    INPUT_TEXT = "C65911"

    NAME_BG = "DDEBF7"            # 材料名称匹配区底色
    NAME_TEXT = "203864"

    CALC_BG = "E2EFDA"            # 自动计算区底色
    CALC_TEXT = "375623"

    SUM_BG = "FCE4D6"             # 汇总行底色
    SUM_TEXT = "C00000"

    TOTAL_ROW_BG = "C00000"       # 全队汇总行底色
    TOTAL_ROW_TEXT = "FFFFFF"

# ==================== 样式对象 ====================
class Styles:
    """预定义样式对象"""

    @staticmethod
    def title_style():
        return {
            "font": Font(size=14, bold=True, color=Colors.TITLE_TEXT),
            "fill": PatternFill(start_color=Colors.TITLE_BG, end_color=Colors.TITLE_BG, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True)
        }

    @staticmethod
    def header_style(bg_color):
        return {
            "font": Font(color="FFFFFF", bold=True, size=10),
            "fill": PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
            "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        }

    @staticmethod
    def input_style():
        return {
            "font": Font(color=Colors.INPUT_TEXT, bold=True, size=10),
            "fill": PatternFill(start_color=Colors.INPUT_BG, end_color=Colors.INPUT_BG, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        }

    @staticmethod
    def name_style():
        return {
            "font": Font(color=Colors.NAME_TEXT, size=10),
            "fill": PatternFill(start_color=Colors.NAME_BG, end_color=Colors.NAME_BG, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        }

    @staticmethod
    def calc_style():
        return {
            "font": Font(color=Colors.CALC_TEXT, size=10),
            "fill": PatternFill(start_color=Colors.CALC_BG, end_color=Colors.CALC_BG, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        }

    @staticmethod
    def sum_style():
        return {
            "font": Font(color=Colors.SUM_TEXT, bold=True, size=10),
            "fill": PatternFill(start_color=Colors.SUM_BG, end_color=Colors.SUM_BG, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        }

    @staticmethod
    def total_row_style():
        return {
            "font": Font(bold=True, size=12, color=Colors.TOTAL_ROW_TEXT),
            "fill": PatternFill(start_color=Colors.TOTAL_ROW_BG, end_color=Colors.TOTAL_ROW_BG, fill_type="solid"),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        }

# ==================== 布局配置 ====================
class Layout:
    """表格布局常量"""
    MAX_CHARACTERS = 10             # 最大角色数
    DATA_START_ROW = 5              # 数据起始行（第5行）
    HEADER_ROW = 4                  # 表头行
    SUM_ROW = 15                    # 汇总行

    # 列定义
    COL_SEQ = 1                     # A: 序号
    COL_CHAR_NAME = 2               # B: 角色名
    COL_CHAR_CURR_LVL = 3           # C: 当前等级
    COL_CHAR_TARGET_LVL = 4         # D: 目标等级
    COL_ARC_NAME = 5                # E: 弧盘名
    COL_ARC_CURR_LVL = 6            # F: 当前弧盘等级
    COL_ARC_TARGET_LVL = 7          # G: 目标弧盘等级
    COL_REMARK = 8                  # H: 备注

    # 角色突破区 (I-T)
    COL_BREAK_MAT_NAME = 9          # I: 突破素材名称
    COL_BREAK_MAT_NUM = 10          # J: 突破素材数量
    COL_LOW_YIXIANG_NAME = 11       # K: 低级异像名称
    COL_LOW_YIXIANG_NUM = 12        # L: 低级异像数量
    COL_MID_YIXIANG_NAME = 13       # M: 中级异像名称
    COL_MID_YIXIANG_NUM = 14        # N: 中级异像数量
    COL_HIGH_YIXIANG_NAME = 15      # O: 高级异像名称
    COL_HIGH_YIXIANG_NUM = 16       # P: 高级异像数量
    COL_EXP_BOOK_LOW = 17           # Q: 新锐攻略
    COL_EXP_BOOK_MID = 18           # R: 资深攻略
    COL_EXP_BOOK_HIGH = 19          # S: 特级攻略
    COL_CHAR_GOLD = 20              # T: 角色突破金币

    # 弧盘强化区 (U-AJ)
    COL_ARC_LOW_MAT_NAME = 21     # U: 低级弧盘素材名称
    COL_ARC_LOW_MAT_NUM = 22      # V: 低级弧盘素材数量
    COL_ARC_MID_MAT_NAME = 23     # W: 中级弧盘素材名称
    COL_ARC_MID_MAT_NUM = 24      # X: 中级弧盘素材数量
    COL_ARC_HIGH_MAT_NAME = 25    # Y: 高级弧盘素材名称
    COL_ARC_HIGH_MAT_NUM = 26     # Z: 高级弧盘素材数量
    COL_ARC_LOW_YX_NAME = 27      # AA: 弧盘低级异像名称
    COL_ARC_LOW_YX_NUM = 28       # AB: 弧盘低级异像数量
    COL_ARC_MID_YX_NAME = 29      # AC: 弧盘中级异像名称
    COL_ARC_MID_YX_NUM = 30       # AD: 弧盘中级异像数量
    COL_ARC_HIGH_YX_NAME = 31     # AE: 弧盘高级异像名称
    COL_ARC_HIGH_YX_NUM = 32      # AF: 弧盘高级异像数量
    COL_DYE_LIGHT = 33            # AG: 淡色染剂
    COL_DYE_COLORLESS = 34        # AH: 无彩染剂
    COL_DYE_CHAOS = 35            # AI: 混沌染剂
    COL_ARC_GOLD = 36             # AJ: 弧盘强化金币

    # 汇总区 (AK-AO)
    COL_TOTAL_GOLD = 37           # AK: 总金币
    COL_TOTAL_DYE_EXP = 38        # AL: 总染剂经验
    COL_TOTAL_LOW_YX = 39         # AM: 总低级异像
    COL_TOTAL_MID_YX = 40         # AN: 总中级异像
    COL_TOTAL_HIGH_YX = 41        # AO: 总高级异像
