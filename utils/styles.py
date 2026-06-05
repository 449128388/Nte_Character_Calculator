"""
Excel 样式工具模块
封装 openpyxl 样式应用，简化代码
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config import Colors

def apply_style(cell, font=None, fill=None, alignment=None, border=None):
    """统一应用样式到单元格"""
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border

def create_header_style(bg_color):
    """创建表头样式"""
    return {
        "font": Font(color="FFFFFF", bold=True, size=10),
        "fill": PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    }

def create_input_style():
    """创建录入区样式"""
    return {
        "font": Font(color=Colors.INPUT_TEXT, bold=True, size=10),
        "fill": PatternFill(start_color=Colors.INPUT_BG, end_color=Colors.INPUT_BG, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    }

def create_name_style():
    """创建材料名称匹配区样式"""
    return {
        "font": Font(color=Colors.NAME_TEXT, size=10),
        "fill": PatternFill(start_color=Colors.NAME_BG, end_color=Colors.NAME_BG, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    }

def create_calc_style():
    """创建自动计算区样式"""
    return {
        "font": Font(color=Colors.CALC_TEXT, size=10),
        "fill": PatternFill(start_color=Colors.CALC_BG, end_color=Colors.CALC_BG, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    }

def create_sum_style():
    """创建汇总区样式"""
    return {
        "font": Font(color=Colors.SUM_TEXT, bold=True, size=10),
        "fill": PatternFill(start_color=Colors.SUM_BG, end_color=Colors.SUM_BG, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    }

def create_total_row_style():
    """创建全队汇总行样式"""
    return {
        "font": Font(bold=True, size=12, color=Colors.TOTAL_ROW_TEXT),
        "fill": PatternFill(start_color=Colors.TOTAL_ROW_BG, end_color=Colors.TOTAL_ROW_BG, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    }

def apply_dict_style(cell, style_dict):
    """从字典应用样式"""
    apply_style(cell, **style_dict)
